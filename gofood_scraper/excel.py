"""Excel output helpers."""

import os

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from gofood_scraper.config import HEADERS


def output_path(output_dir, area, district, category_link, suffix=""):
    district_path = "-{}".format(district) if district else ""
    return os.path.join(output_dir, "{}{}-{}{}.xlsx".format(area, district_path, category_link, suffix))


def column(header):
    """Return the column letter for a header name from config.HEADERS, e.g. column("Address") -> "I"."""
    return get_column_letter(HEADERS.index(header) + 1)


def set_headers(sheet_obj):
    for i, header in enumerate(HEADERS, start=1):
        sheet_obj["{}1".format(get_column_letter(i))] = header


def new_results_workbook():
    wb_obj = Workbook()
    sheet_obj = wb_obj.active
    sheet_obj.title = "results"
    set_headers(sheet_obj)
    return wb_obj, sheet_obj
